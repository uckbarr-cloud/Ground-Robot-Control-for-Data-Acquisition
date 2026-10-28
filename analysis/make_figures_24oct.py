"""
make_figures_24oct.py - figures 7 to 12 from the condition reports (24 Oct, plus 21 Oct for 100 cm).

Reads every Condition_Report_20261021_*.docx and Condition_Report_20261024_*.docx under reports/,
so every number in the figures comes straight from the reports. Writes six PNG figures and two
CSV files to figures_24oct/ and prints a summary.

Run from the project folder:  python3 analysis/make_figures_24oct.py
"""
import csv, glob, os, re
from collections import OrderedDict
import numpy as np
import cv2
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from docx import Document

REPORTS_DIR = "reports"
OUT_DIR = "figures_24oct"

# Times used for the break-even figure (fig10)
MANUAL_MIN = 17 + 37 / 60   # manual inspection of the same 5 m, from the photo times (17 min 37 s)
REVIEW_S = 45               # timed review of report 144211, about 45 s per finding

# Scale that is actually true at each standoff (card scale bar 23 Oct; 1.27 measured 6 Oct at 100 cm).
SCALE_40, SCALE_60, SCALE_100 = 0.50, 0.77, 1.27
RUN_60CM = "20261024_145854"     # its report was generated with 0.50, so widths are rescaled from pixels
# 21 Oct run 110823 drifted from 100 cm to about 30 cm by its last stops, so that finding has no known scale
EXCLUDE_WIDTH = {("20261021_110823", 23)}

SETUP = OrderedDict([
    ("20261024_131254", "40 cm, test sheets"), ("20261024_132539", "40 cm, test sheets"),
    ("20261024_133911", "40 cm, test sheets"), ("20261024_141022", "40 cm, clean wall"),
    ("20261024_142236", "40 cm, clean wall"), ("20261024_142923", "40 cm, clean wall"),
    ("20261024_144211", "40 cm, clean wall"), ("20261024_145854", "60 cm, clean wall")])
NOTES = {"20261024_133911": "natural re-shoot", "20261024_142236": "stopped, Wi-Fi drop",
         "20261024_144211": "re-shoot, lens covered"}
# Stop 21 deck edge, identified by eye in the annotated photos (in 144211, stop 21 is joint lines instead)
DECK_EDGE_RUNS = ["20261024_131254", "20261024_132539", "20261024_133911",
                  "20261024_141022", "20261024_142923"]
# Photo panel: run, finding, label, crop (x0, y0, x1, y1 as fractions of the annotated photo)
PANEL = [("20261024_144211", "F-11", "(a) Horizontal joints seen at a steep angle", (0.40, 0.00, 1.00, 0.62)),
         ("20261024_144211", "F-13", "(b) Vertical joint near the floor", (0.22, 0.15, 0.82, 0.77)),
         ("20261024_132539", "F-08", "(c) Deck edge against the bright exit", (0.40, 0.00, 1.00, 0.62))]

COL = {"40 cm, test sheets": "#E69F00", "40 cm, clean wall": "#0072B2",
       "60 cm, clean wall": "#009E73", "100 cm (21 Oct)": "#CC79A7"}
plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False,
                     "savefig.dpi": 300, "savefig.bbox": "tight"})


# ------------------------------------------------------------------ parsing
def cell(c):
    return " ".join(p.text.strip() for p in c.paragraphs).strip()


def num(s):
    try:
        return float(str(s).replace("%", "").strip())
    except ValueError:
        return None


def hhmm(s):
    h, m = s.split(":")
    return int(h) * 60 + int(m)


def parse_report(path):
    doc = Document(path)
    body = "\n".join(p.text for p in doc.paragraphs)
    kv = {}
    for t in doc.tables:
        for row in t.rows:
            c = [cell(x) for x in row.cells]
            if len(c) >= 2 and c[0] and c[1] and len(c[0]) < 40:
                kv.setdefault(c[0], c[1])
    run = re.search(r"2026\d{4}_\d{6}", kv.get("Run ID", "") + " " + os.path.basename(path)).group(0)
    win = re.findall(r"\d{1,2}:\d{2}", kv.get("Capture window", ""))
    gen = re.findall(r"\d{1,2}:\d{2}", kv.get("Report generated", ""))
    same_day = kv.get("Inspection date", "?") in kv.get("Report generated", "")
    stops = re.match(r"(\d+) of (\d+)", kv.get("Waypoints photographed", "0 of 0"))
    spacing = re.search(r"([0-9.]+) m spacing", kv.get("Route", ""))
    scale = re.search(r"wall ([0-9.]+) mm/px", body)
    r = dict(run=run, date=run[:8], path=path, status=kv.get("Run status", ""),
             start=hhmm(win[0]) if win else None, end=hhmm(win[1]) if len(win) > 1 else None,
             generated=hhmm(gen[-1]) if gen else None,
             stops=int(stops.group(1)), planned=int(stops.group(2)),
             photos=int(re.match(r"\d+", kv.get("Images analysed", "0")).group(0)),
             reshoots=int(num(kv.get("Reshoots triggered", "0")) or 0),
             spacing=float(spacing.group(1)) if spacing else 0.2,
             report_scale=float(scale.group(1)) if scale else None,
             overall=kv.get("OVERALL INDICATIVE CONDITION", ""), findings=[])
    for t in doc.tables:
        hdr = [cell(x) for x in t.rows[0].cells]
        if not hdr or hdr[0] != "ID" or "Severity" not in " ".join(hdr):
            continue
        for row in t.rows[1:]:
            c = [cell(x) for x in row.cells]
            if not c[0].startswith("F-"):
                continue
            r["findings"].append(dict(id=c[0], stop=int(re.search(r"WP (\d+)", c[1]).group(1)),
                                      surface=c[2], width=num(c[4]), length=num(c[5]),
                                      conf=num(c[6]), severity=c[7]))
    # a report regenerated later (new date, or long after the run) has no on-site report time
    if r["generated"] is not None and (not same_day or not 0 <= r["generated"] - r["end"] <= 10):
        r["generated"] = None
    r["route_m"] = round((r["stops"] - 1) * r["spacing"], 2)
    r["complete"] = r["stops"] == r["planned"]
    if run.startswith("20261021"):
        r["setup"], r["true_scale"] = "100 cm (21 Oct)", SCALE_100
    else:
        r["setup"] = SETUP.get(run, "unknown")
        r["true_scale"] = SCALE_60 if run == RUN_60CM else SCALE_40
    for f in r["findings"]:
        f["width_px"] = f["width"] / r["report_scale"] if f["width"] and r["report_scale"] else None
        f["width_true"] = f["width_px"] * r["true_scale"] if f["width_px"] else None
        f["length_true"] = f["length"] / r["report_scale"] * r["true_scale"] if f["length"] else None
    return r


def finding_image(path, fid):
    """The annotated photo the report shows for one finding, as an RGB array."""
    doc = Document(path)
    current = None
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            m = re.match(r"(F-\d+)\s+\|", "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t")))
            current = m.group(1) if m else current
        elif tag == "tbl" and current == fid:
            for b in el.iter():
                if b.tag.endswith("}blip"):
                    rid = b.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    img = cv2.imdecode(np.frombuffer(doc.part.related_parts[rid].blob, np.uint8), cv2.IMREAD_COLOR)
                    return cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
    return None


def load_reports():
    newest = {}
    for p in glob.glob(os.path.join(REPORTS_DIR, "**", "*Condition_Report_2026102[14]_*.docx"), recursive=True):
        if os.path.basename(p).startswith("~$"):
            continue
        run = re.search(r"2026102[14]_\d{6}", os.path.basename(p)).group(0)
        if run not in newest or os.path.getmtime(p) > os.path.getmtime(newest[run]):
            newest[run] = p
    return {run: parse_report(p) for run, p in sorted(newest.items())}


def clock(m):
    return "" if m is None else "%02d:%02d" % (m // 60, m % 60)


# ------------------------------------------------------------------ figures
def fig7_width_vs_scale(R):
    pts = []
    for r in R.values():
        for f in r["findings"]:
            if f["surface"] == "Wall" and f["width_true"] and (r["run"], f["stop"]) not in EXCLUDE_WIDTH:
                group = "40 cm" if r["true_scale"] == SCALE_40 else ("60 cm" if r["true_scale"] == SCALE_60 else "100 cm")
                pts.append((r["true_scale"], f["width_true"], group))
    x = np.array([p[0] for p in pts]); y = np.array([p[1] for p in pts])
    k = (x * y).sum() / (x * x).sum()                          # width = k pixels x scale, through the origin
    r2 = 1 - ((y - k * x) ** 2).sum() / ((y - y.mean()) ** 2).sum()
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(6.3, 4.2))
    colours = {"40 cm": COL["40 cm, clean wall"], "60 cm": COL["60 cm, clean wall"], "100 cm": COL["100 cm (21 Oct)"]}
    stats = {}
    for g in ["40 cm", "60 cm", "100 cm"]:
        sel = np.array([p[2] == g for p in pts])
        if not sel.any():
            continue
        xs, ys = x[sel], y[sel]
        ax.scatter(xs * np.exp(rng.normal(0, 0.035, sel.sum())), ys, s=12, alpha=0.45,
                   color=colours[g], edgecolor="none", label="%s, %d findings" % (g, sel.sum()))
        q1, med, q3 = np.percentile(ys, [25, 50, 75])
        ax.errorbar(xs[0], med, yerr=[[med - q1], [q3 - med]], fmt="o", ms=7, color="black", capsize=3, zorder=5)
        stats[g] = (xs[0], sel.sum(), med, med / xs[0])
    xx = np.linspace(0.38, 1.7, 50)
    ax.plot(xx, k * xx, color="black", lw=1.2, label="fit: %.0f pixels wide at every standoff (R$^2$ = %.2f)" % (k, r2))
    ax.plot(xx, xx, color="grey", lw=0.9, ls=":", label="one pixel")
    for yv, txt, xpos, ha in [(3, "report rule: High above 3 mm", 0.39, "left"),
                              (1, "report rule: Low below 1 mm", 0.39, "left"),
                              (0.3, "Eurocode 2 crack limit, 0.3 mm", 1.68, "right")]:
        ax.axhline(yv, color="#B22222", lw=0.9, ls="--")
        ax.text(xpos, yv * 1.08, txt, ha=ha, va="bottom", fontsize=7.5, color="#B22222")
    ax.set_xscale("log"); ax.set_yscale("log")
    ax.set_xlim(0.38, 1.7); ax.set_ylim(0.1, 100)
    ax.set_xticks([SCALE_40, SCALE_60, SCALE_100])
    ax.set_xticklabels(["%.2f\n(40 cm)" % SCALE_40, "%.2f\n(60 cm)" % SCALE_60, "%.2f\n(100 cm)" % SCALE_100])
    ax.minorticks_off()
    ax.set_yticks([0.1, 0.3, 1, 3, 10, 30, 100]); ax.set_yticklabels(["0.1", "0.3", "1", "3", "10", "30", "100"])
    ax.set_xlabel("Pixel size on the wall, mm per pixel (standoff)")
    ax.set_ylabel("Measured mean width (mm)")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=2, fontsize=7.5, frameon=False)
    fig.savefig(os.path.join(OUT_DIR, "fig7_width_vs_pixel_size.png")); plt.close(fig)
    return k, r2, len(pts), stats


def fig8_threshold(R24):
    th = np.arange(25, 96)
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(6.3, 4.6), sharex=True, gridspec_kw={"height_ratios": [2.2, 1]})
    rows = {}
    for setup in ["40 cm, test sheets", "40 cm, clean wall", "60 cm, clean wall"]:
        runs = [r for r in R24.values() if r["setup"] == setup and r["complete"]]
        conf = np.array([f["conf"] for r in runs for f in r["findings"]])
        metres = sum(r["route_m"] for r in runs)
        per_m = np.array([(conf >= t).sum() / metres for t in th])
        rows[setup] = (len(runs), metres, per_m)
        a1.step(th, per_m, where="post", color=COL[setup], lw=1.6,
                label="%s (%d run%s, %.0f m)" % (setup, len(runs), "s" if len(runs) > 1 else "", metres))
    flagged = np.array([sum(any(f["conf"] >= t for f in r["findings"]) for r in R24.values()) for t in th])
    a2.step(th, flagged, where="post", color="black", lw=1.4)
    for a in (a1, a2):
        a.axvline(25, color="grey", ls="--", lw=0.8)
    a1.text(25.6, 0.35, "threshold used, 25%", rotation=90, fontsize=7.5, color="grey", va="bottom", ha="left")
    a2.set_xlim(24, 96)
    a1.set_ylabel("False calls per metre")
    a1.legend(frameon=False, fontsize=7.5)
    a2.set_ylabel("Reports still\n'significant defects'")
    a2.set_ylim(-0.3, len(R24) + 0.5); a2.set_yticks(range(0, len(R24) + 1, 2))
    a2.set_xlabel("Confidence threshold (%)")
    fig.savefig(os.path.join(OUT_DIR, "fig8_false_calls_vs_threshold.png")); plt.close(fig)
    return th, flagged, rows


def fig9_run_stop_map(R24):
    runs = list(R24.values())
    n_stops = max(r["planned"] for r in runs)
    M = np.full((len(runs), n_stops), -1)
    for i, r in enumerate(runs):
        M[i, :r["stops"]] = 0
        for f in r["findings"]:
            M[i, f["stop"]] = 1
    core = [i for i, r in enumerate(runs) if r["setup"].startswith("40 cm") and r["complete"]]
    freq = (M[core] == 1).sum(0)
    hot = [s for s in range(n_stops) if freq[s] == len(core)]
    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    for i, r in enumerate(runs):
        for s in range(n_stops):
            v = M[i, s]
            ax.add_patch(Rectangle((s, i), 0.92, 0.86, facecolor={1: COL[r["setup"]], 0: "#EEEEEE", -1: "white"}[v],
                                   edgecolor="#BBBBBB" if v == -1 else "none", hatch="////" if v == -1 else None, lw=0.3))
        ax.text(n_stops + 0.3, i + 0.43, "%d" % (M[i] == 1).sum(), va="center", fontsize=7.5)
    blocks, cur = [], [core[0]]
    for i in core[1:]:
        if i == cur[-1] + 1:
            cur.append(i)
        else:
            blocks.append(cur); cur = [i]
    blocks.append(cur)
    for s in hot:
        for b in blocks:
            ax.add_patch(Rectangle((s - 0.08, b[0] - 0.08), 1.08, len(b) - 0.06,
                                   fill=False, edgecolor="#B22222", lw=1.3))
    for s in range(n_stops):
        ax.text(s + 0.46, len(runs) + 0.35, "%d" % freq[s], ha="center", fontsize=7,
                color="#B22222" if s in hot else "black", fontweight="bold" if s in hot else "normal")
    ax.text(-0.3, len(runs) + 0.35, "of %d complete 40 cm runs" % len(core), ha="right", fontsize=7)
    ax.text(n_stops + 0.3, -0.45, "flags", fontsize=7)
    ax.set_xlim(-0.2, n_stops + 1.5); ax.set_ylim(len(runs) + 0.9, -0.7)
    ax.set_yticks([i + 0.43 for i in range(len(runs))])
    ax.set_yticklabels(["%s  %s%s" % (r["run"][-6:], r["setup"], ("  (" + NOTES[r["run"]] + ")") if r["run"] in NOTES else "")
                        for r in runs], fontsize=7.5)
    ax.set_xticks([s + 0.46 for s in range(0, n_stops, 5)]); ax.set_xticklabels(range(0, n_stops, 5))
    ax.set_xlabel("Stop number (0.2 m apart; stop 25 = 5.0 m)")
    ax.tick_params(length=0)
    for side in ("left", "bottom"):
        ax.spines[side].set_visible(False)
    fig.savefig(os.path.join(OUT_DIR, "fig9_run_by_stop_map.png")); plt.close(fig)
    return M, core, freq, hot


def stats_repeatability(M, core, iters=20000):
    X = (M[core] == 1).astype(int)
    n, N = X.shape
    k = X.sum(0)
    p1 = X.mean()
    p_obs = np.mean((k * (k - 1) + (n - k) * (n - k - 1)) / (n * (n - 1)))
    p_exp = p1 ** 2 + (1 - p1) ** 2
    kappa = (p_obs - p_exp) / (1 - p_exp)
    obs_all = (k == n).sum()
    rng = np.random.default_rng(1)
    hits = sum(((np.array([rng.permutation(row) for row in X]).sum(0) == n).sum() >= obs_all) for _ in range(iters))
    return kappa, obs_all, hits / iters


def time_components(R24):
    runs = sorted(R24.values(), key=lambda r: r["start"])
    capture = [r["end"] - r["start"] for r in runs if r["complete"] and r["reshoots"] == 0
               and r["setup"].startswith("40 cm") and "Completed" in r["status"]]
    latency = [r["generated"] - r["end"] for r in runs if r["generated"] is not None]
    gaps = [b["start"] - a["end"] for a, b in zip(runs, runs[1:])]
    reset = [g for g in gaps if g <= 10]
    return float(np.median(capture)), float(np.median(latency)), float(np.median(reset)), capture, latency, gaps


def fig10_break_even(R24):
    cap, lat, reset, *_ = time_components(R24)
    fixed = cap + lat + reset
    counts = [len(r["findings"]) for r in R24.values() if r["setup"].startswith("40 cm") and r["complete"]]
    ns = [min(counts), float(np.median(counts)), max(counts)]
    x = np.linspace(0, 60, 121)
    fig, ax = plt.subplots(figsize=(6.3, 4.0))
    ax.fill_between(x, fixed + ns[0] * x / 60, fixed + ns[2] * x / 60, color=COL["40 cm, clean wall"], alpha=0.15)
    out = []
    for n, ls, word in zip(ns, [":", "-", "--"], ["fewest", "median", "most"]):
        be = (MANUAL_MIN - fixed) * 60 / n
        out.append((n, be))
        tag = "break-even %.0f s" % be if be > 0 else "no break-even"
        ax.plot(x, fixed + n * x / 60, color=COL["40 cm, clean wall"], ls=ls, lw=1.4,
                label="robot, %g findings to review (%s; %s)" % (n, word, tag))
        if 0 < be <= 60:
            ax.plot(be, MANUAL_MIN, "o", color="black", ms=4, zorder=5)
    ax.axhline(MANUAL_MIN, color="black", lw=1.4, label="manual, %d min %02d s" % (int(MANUAL_MIN), round(MANUAL_MIN % 1 * 60)))
    if REVIEW_S is not None:
        ax.axvline(REVIEW_S, color="#B22222", ls="--", lw=1, label="measured review time, %g s" % REVIEW_S)
    ax.text(1, fixed - 0.5, "robot before review: capture %.0f + report %.0f + setup %.0f = %.0f min" % (cap, lat, reset, fixed),
            fontsize=7.5, va="top")
    ax.set_xlim(0, 60); ax.set_ylim(0, max(fixed + ns[2], MANUAL_MIN) + 3)
    ax.set_xlabel("Review time per finding (s)"); ax.set_ylabel("Time per 5 m of wall (min)")
    ax.legend(frameon=False, fontsize=7.5, loc="upper left")
    fig.savefig(os.path.join(OUT_DIR, "fig10_review_break_even.png")); plt.close(fig)
    return fixed, (cap, lat, reset), ns, out


def fig11_panel(R):
    fig, axes = plt.subplots(1, len(PANEL), figsize=(9.0, 2.6))
    for ax, (run, fid, label, (x0, y0, x1, y1)) in zip(axes, PANEL):
        img = finding_image(R[run]["path"], fid)
        f = next(f for f in R[run]["findings"] if f["id"] == fid)
        h, w = img.shape[:2]
        ax.imshow(img[int(y0 * h):int(y1 * h), int(x0 * w):int(x1 * w)])
        ax.set_xticks([]); ax.set_yticks([])
        for s in ax.spines.values():
            s.set_visible(True); s.set_linewidth(0.5)
        ax.set_title("%s\nrun %s, %s, stop %d, confidence %.0f%%" % (label, run[-6:], fid, f["stop"], f["conf"]),
                     fontsize=7.5, loc="left")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT_DIR, "fig11_false_alarm_types.png")); plt.close(fig)


def fig12_timeline(R24):
    runs = sorted(R24.values(), key=lambda r: r["start"])
    fig, ax = plt.subplots(figsize=(7.2, 3.4))
    labels = []
    rows = []
    for r in runs:
        rows.append(("run", r))
    # the gap between the sheet runs and the clean runs holds the manual inspection and taking the sheets down
    i_gap = next(i for i, (a, b) in enumerate(zip(runs, runs[1:])) if b["start"] - a["end"] > 10)
    rows.insert(i_gap + 1, ("gap", (runs[i_gap]["end"], runs[i_gap + 1]["start"])))
    for y, (kind, obj) in enumerate(rows):
        if kind == "gap":
            a, b = obj
            ax.barh(y, b - a, left=a, height=0.6, color="#DDDDDD", edgecolor="#999999", lw=0.5)
            ax.text(b + 1, y, "manual inspection and sheets taken down, within this %d min gap" % (b - a),
                    va="center", fontsize=7)
            labels.append("manual")
            continue
        r = obj
        stopped = not r["complete"]
        ax.barh(y, r["end"] - r["start"], left=r["start"], height=0.6, color=COL[r["setup"]],
                hatch="////" if stopped else None, edgecolor="white" if stopped else "none")
        note = ", " + NOTES[r["run"]] if r["run"] in NOTES else ""
        if r["generated"] is not None:
            ax.plot(r["generated"], y, marker="D", ms=3.5, color="black")
        else:
            note += ", report regenerated later"
        ax.text((r["generated"] or r["end"]) + 2, y, "%d stops, %d flags%s" % (r["stops"], len(r["findings"]), note),
                va="center", fontsize=7)
        labels.append(r["run"][-6:])
    t0 = (min(r["start"] for r in runs) // 10) * 10
    t1 = max(r["generated"] or r["end"] for r in runs) + 40
    ax.set_xlim(t0, t1)
    ax.set_xticks(range(t0, t1, 10)); ax.set_xticklabels([clock(t) for t in range(t0, t1, 10)], fontsize=7.5)
    ax.set_yticks(range(len(rows))); ax.set_yticklabels(labels, fontsize=7.5)
    ax.invert_yaxis()
    ax.set_xlabel("Time on 24 Oct (bars: capture window from the reports; diamond: draft report ready)")
    handles = [Rectangle((0, 0), 1, 1, color=COL[s]) for s in ["40 cm, test sheets", "40 cm, clean wall", "60 cm, clean wall"]]
    ax.legend(handles, ["40 cm, test sheets", "40 cm, clean wall", "60 cm, clean wall"], frameon=False,
              fontsize=7.5, loc="upper right")
    fig.savefig(os.path.join(OUT_DIR, "fig12_field_day_timeline.png")); plt.close(fig)


# ------------------------------------------------------------------ main
def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    R = load_reports()
    R24 = OrderedDict((run, R[run]) for run in SETUP if run in R)
    missing = [run for run in SETUP if run not in R]
    if missing:
        print("Missing 24 Oct reports:", missing)
    n21 = sum(1 for run in R if run.startswith("20261021"))
    print("Reports read: %d from 24 Oct, %d from 21 Oct" % (len(R24), n21))

    with open(os.path.join(OUT_DIR, "runs.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["run", "setup", "status", "capture_start", "capture_end", "capture_min", "report_ready",
                    "stops", "planned", "route_m", "photos", "findings", "per_m", "reshoots", "report_scale",
                    "true_scale", "overall"])
        for r in R.values():
            w.writerow([r["run"], r["setup"], r["status"], clock(r["start"]), clock(r["end"]), r["end"] - r["start"],
                        clock(r["generated"]), r["stops"], r["planned"], r["route_m"], r["photos"], len(r["findings"]),
                        round(len(r["findings"]) / r["route_m"], 2) if r["route_m"] else "", r["reshoots"],
                        r["report_scale"], r["true_scale"], r["overall"]])
    with open(os.path.join(OUT_DIR, "findings.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["run", "setup", "id", "stop", "surface", "width_mm_report", "width_px", "width_mm_true",
                    "length_mm_true", "confidence", "severity"])
        for r in R.values():
            for f in r["findings"]:
                w.writerow([r["run"], r["setup"], f["id"], f["stop"], f["surface"], f["width"],
                            round(f["width_px"], 1) if f["width_px"] else "", round(f["width_true"], 1) if f["width_true"] else "",
                            round(f["length_true"]) if f["length_true"] else "", f["conf"], f["severity"]])

    k, r2, npts, wstats = fig7_width_vs_scale(R)
    th, flagged, rows = fig8_threshold(R24)
    M, core, freq, hot = fig9_run_stop_map(R24)
    kappa, n_all, p_all = stats_repeatability(M, core)
    fixed, parts, ns, be = fig10_break_even(R24)
    fig11_panel(R)
    fig12_timeline(R24)

    allf = [f for r in R24.values() for f in r["findings"]]
    metres = sum(r["route_m"] for r in R24.values())
    print("\n24 Oct: %d photos, %d findings over %.1f m (%.2f per m); severities %s; overall %s"
          % (sum(r["photos"] for r in R24.values()), len(allf), metres, len(allf) / metres,
             sorted({f["severity"] for f in allf}), sorted({r["overall"] for r in R24.values()})))
    print("Both days: %d reports, %d findings, severities %s"
          % (len(R), sum(len(r["findings"]) for r in R.values()), sorted({f["severity"] for r in R.values() for f in r["findings"]})))
    print("\nRun     setup                 capture      min  stops  photos  flags  per_m  reshoots")
    for r in R24.values():
        print("%s  %-20s  %s-%s  %3d  %2d/%2d  %5d  %5d  %5.2f  %d" % (r["run"][-6:], r["setup"], clock(r["start"]), clock(r["end"]),
              r["end"] - r["start"], r["stops"], r["planned"], r["photos"], len(r["findings"]),
              len(r["findings"]) / r["route_m"], r["reshoots"]))
    print("\nFig 7: width = %.1f px x scale, R2 = %.2f, n = %d (excluded %s)" % (k, r2, npts, sorted(EXCLUDE_WIDTH)))
    for g, (sc, n, med, px) in wstats.items():
        print("   %-6s %.2f mm/px  n=%3d  median %.1f mm = %.0f px" % (g, sc, n, med, px))
    print("Fig 8: false calls per m at 25/50/70/90%:")
    for setup, (nr, m, per_m) in rows.items():
        print("   %-20s %s" % (setup, "  ".join("%.2f" % per_m[t - 25] for t in (25, 50, 70, 90))))
    print("   reports still 'significant defects' at 25/50/70/80/90%%: %s" % [int(flagged[t - 25]) for t in (25, 50, 70, 80, 90)])
    print("Fig 9: stops flagged in all %d complete 40 cm runs: %s (shuffle test p = %.3f); Fleiss kappa %.2f"
          % (len(core), hot, p_all, kappa))
    e = [f for run in DECK_EDGE_RUNS if run in R for f in R[run]["findings"] if f["stop"] == 21]
    for key, name in [("width", "width mm"), ("length", "length mm"), ("conf", "confidence %")]:
        v = np.array([f[key] for f in e])
        print("   stop 21 deck edge, %d runs, %-12s mean %.1f  SD %.1f  CV %.1f%%" % (len(v), name, v.mean(), v.std(ddof=1), 100 * v.std(ddof=1) / v.mean()))
    cap, lat, reset = parts
    print("Fig 10: robot before review %.0f min (capture %.0f, report %.0f, setup %.0f); manual %.2f min" % (fixed, cap, lat, reset, MANUAL_MIN))
    for n, b in be:
        print("   %g findings: break-even review time %s" % (n, "%.0f s per finding" % b if b > 0 else "none, robot slower even with instant review"))
    print("\nFigures and CSV files saved in %s/" % OUT_DIR)


if __name__ == "__main__":
    main()
